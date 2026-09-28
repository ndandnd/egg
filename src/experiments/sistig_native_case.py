"""Source-faithful, no-solver adapter for the public Sistig Hildenbrand case.

This script reads the DOI-pinned workbook bundle as data only. It preserves all
37 service rows and the full directed 5x5 deadhead matrix, then creates two
separate single-depot scenario variants for the native recharge schema. The
variants are explicit restrictions of the source two-depot network, not a
reproduction of the publisher's two-depot optimization.

Requires numpy, openpyxl, and scipy only for source intake/regeneration; these
are imported lazily. It never extracts the large source archive or invokes the
optimizer. NativeCase validation, interval compilation, and witness replay are
pure and run only when requested with --validate-native.
"""
from __future__ import annotations

import argparse
from dataclasses import asdict
import hashlib
import io
import json
import math
from pathlib import Path, PurePosixPath
import sys
import zipfile

PROJECT_ROOT = Path(__file__).resolve().parents[2]
WORKSPACE_ROOT = PROJECT_ROOT.parent
if str(PROJECT_ROOT / "src") not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT / "src"))
DEFAULT_ARCHIVE = WORKSPACE_ROOT / "research-20260927/agent-notes/public-data/sistig-26088190-v1.zip"
DEFAULT_OUTPUT = PROJECT_ROOT / "data/public/sistig_26088190_v1/hildenbrand_native_cases.json"
EBERBACH_OUTPUT = PROJECT_ROOT / "data/public/sistig_26088190_v1/eberbach_native_case.json"
DOI = "10.6084/m9.figshare.26088190.v1"
FIGSHARE_URL = "https://figshare.com/articles/dataset/Dataset_for_Evaluating_Costs_and_Operations_of_Public_Bus_Fleet_Electrification/26088190"
PAPER_URL = "https://www.nature.com/articles/s44333-025-00030-y"
LICENSE_URL = "https://creativecommons.org/licenses/by/4.0/"
ZIP_SHA256 = "3c6c2ed7c45fc441f4cb6108d830769b60c0de5e44e1a3a6fc72bed4f5f07317"
ZIP_MD5 = "89b04ee5e8cd54997df51211d2c264d2"
ZIP_BYTES = 267_312_231
OPERATOR = "10__data/20__Selected_transport_operators/020__Hildenbrand/"
OPERATOR_CONFIGS = {
    "hildenbrand": {
        "member_prefix": OPERATOR, "directory": "020__Hildenbrand",
        "name": "Hildenbrand", "trip_prefix": "H", "service_count": 37,
        "depot_ids": (15, 16), "max_vehicles": 37,
        "case_prefix": "sistig_hildenbrand_single_depot",
    },
    "eberbach": {
        "member_prefix": "10__data/20__Selected_transport_operators/003__Stadtwerke Eberbach/",
        "directory": "003__Stadtwerke Eberbach", "name": "Stadtwerke Eberbach",
        "trip_prefix": "E", "service_count": 105,
        "depot_ids": (36,), "max_vehicles": 105,
        "case_prefix": "sistig_eberbach_single_depot",
    },
}
EXPECTED_MEMBERS = (
    "bus_route_info.xlsx",
    "bus_stops.xlsx",
    "deadhead_trip_matrix.mat",
    "itineraries.xlsx",
    "itineraries_course.mat",
    "trip_set.xlsx",
)

# Explicit EGG baseline scenario choices, taken from the research lead's
# selected protocol and the paper's Table 2. These are assumptions, never
# source-observed vehicle telemetry or charger deployment.
EB3 = {
    "paper_vehicle_id": "4103",
    "paper_vehicle_name": "EB-3",
    "usable_battery_kwh": 400.0,
    "installed_battery_kwh_reference_only": 600.0,
    "service_traction_kwh_per_km": 1.20,
    "deadhead_traction_kwh_per_km": 0.96,
    "auxiliary_kwh_per_hour": 16.0,
    "charging_power_kw_usable": 360.0,
    "charging_power_kw_installed_reference_only": 450.0,
}
HORIZON_START_SEC = 0
HORIZON_END_SEC = 30 * 60 * 60
VEHICLE_COST_SYNTHETIC = 100.0
DEADHEAD_COST_PER_MIN_SYNTHETIC = 0.0


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for b in iter(lambda: f.read(1024 * 1024), b""):
            h.update(b)
    return h.hexdigest()


def md5_file(path: Path) -> str:
    h = hashlib.md5()
    with path.open("rb") as f:
        for b in iter(lambda: f.read(1024 * 1024), b""):
            h.update(b)
    return h.hexdigest()


def parse_day_time(value: str) -> int:
    """Parse dd:HH:MM:SS exactly to seconds since selected-day boundary."""
    if not isinstance(value, str):
        raise ValueError(f"Expected source dd:HH:MM:SS text, got {value!r}")
    parts = value.split(":")
    if len(parts) != 4 or not all(p.isdigit() for p in parts):
        raise ValueError(f"Malformed source dd:HH:MM:SS value: {value!r}")
    days, hours, minutes, seconds = map(int, parts)
    if hours >= 24 or minutes >= 60 or seconds >= 60:
        raise ValueError(f"Out-of-range source time: {value!r}")
    return days * 86400 + hours * 3600 + minutes * 60 + seconds


def parse_duration(value: str) -> int:
    """Parse source HH:MM:SS duration to seconds (hours may exceed 23)."""
    if not isinstance(value, str):
        raise ValueError(f"Expected source HH:MM:SS duration, got {value!r}")
    parts = value.split(":")
    if len(parts) != 3 or not all(p.isdigit() for p in parts):
        raise ValueError(f"Malformed source HH:MM:SS duration: {value!r}")
    hours, minutes, seconds = map(int, parts)
    if minutes >= 60 or seconds >= 60:
        raise ValueError(f"Out-of-range source duration: {value!r}")
    return hours * 3600 + minutes * 60 + seconds


def exact_native_minutes(seconds: int) -> int | float:
    """Convert seconds to exact (whole or half) minutes without rounding."""
    if isinstance(seconds, bool) or not isinstance(seconds, int) or seconds % 30:
        raise ValueError(f"Native half-minute grid cannot represent {seconds!r} seconds exactly")
    if seconds % 60 == 0:
        return seconds // 60
    return seconds / 60.0  # half values are exact in binary floating point


def read_xlsx(zf: zipfile.ZipFile, member: str) -> tuple[list[str], list[dict]]:
    try:
        import openpyxl
    except ImportError as exc:  # artifact loading and tests need not install intake dependencies
        raise RuntimeError("Regenerating source data requires openpyxl") from exc
    workbook = openpyxl.load_workbook(io.BytesIO(zf.read(member)), read_only=True,
                                      data_only=True, keep_links=False)
    try:
        sheet = workbook[workbook.sheetnames[0]]
        rows = sheet.iter_rows(values_only=True)
        header = [str(x) if x is not None else "" for x in next(rows)]
        records = []
        for row in rows:
            if not any(x is not None for x in row):
                continue
            records.append({header[i]: row[i] if i < len(row) else None for i in range(len(header))})
        return header, records
    finally:
        workbook.close()


def member_provenance(zf: zipfile.ZipFile, archive_path: Path,
                      operator: str = "hildenbrand") -> dict:
    prefix = OPERATOR_CONFIGS[operator]["member_prefix"]
    files = []
    for short in EXPECTED_MEMBERS:
        name = prefix + short
        info = zf.getinfo(name)
        payload = zf.read(name)
        files.append({
            "member_path": name,
            "uncompressed_bytes": info.file_size,
            "crc32_hex": f"{info.CRC:08x}",
            "sha256": sha256_bytes(payload),
            "used_as": {
                "bus_route_info.xlsx": "route labels and declared trip totals",
                "bus_stops.xlsx": "place identifiers, names and depot flags",
                "deadhead_trip_matrix.mat": "directed OD distance/time/elevation fields",
                "itineraries.xlsx": "service itinerary identifiers and endpoints",
                "itineraries_course.mat": "fingerprinted for complete source provenance; opaque MCOS route-course table, not parsed or used by endpoint model",
                "trip_set.xlsx": f"the full {OPERATOR_CONFIGS[operator]['service_count']} mandatory fixed service records",
            }[short],
        })
    return {
        "archive_bytes": archive_path.stat().st_size,
        "archive_sha256": sha256_file(archive_path),
        "archive_md5": md5_file(archive_path),
        # Figshare supplied the MD5. SHA-256 is our intake fingerprint, not a
        # publisher-supplied checksum; the exact values are also recorded in
        # the verified local intake receipt.
        "archive_verified_intake_sha256": ZIP_SHA256,
        "archive_figshare_supplied_md5": ZIP_MD5,
        "archive_verified_intake_bytes": ZIP_BYTES,
        "source_archive_bytes_included": False,
        "archive_stored_outside_project_data_directory": True,
        "member_provenance": files,
    }


def extract_source(zf: zipfile.ZipFile, archive_path: Path,
                   operator: str = "hildenbrand") -> dict:
    config = OPERATOR_CONFIGS[operator]
    prefix = config["member_prefix"]
    try:
        import numpy as np
        from scipy.io import loadmat
    except ImportError as exc:
        raise RuntimeError("Regenerating source data requires numpy and scipy") from exc
    paths = [PurePosixPath(x.filename) for x in zf.infolist()]
    if any(p.is_absolute() or ".." in p.parts or "\\" in str(p) for p in paths):
        raise ValueError("Unsafe ZIP member path found")
    if len(zf.namelist()) != len(set(zf.namelist())):
        raise ValueError("Duplicate ZIP member paths")
    _, trips = read_xlsx(zf, prefix + "trip_set.xlsx")
    _, stops = read_xlsx(zf, prefix + "bus_stops.xlsx")
    _, routes = read_xlsx(zf, prefix + "bus_route_info.xlsx")
    _, itineraries = read_xlsx(zf, prefix + "itineraries.xlsx")
    if len(trips) != config["service_count"]:
        raise ValueError(f"Pinned {config['name']} service count changed: expected {config['service_count']}, got {len(trips)}")
    trip_fields = {
        "trip_id", "day_type_id", "bus_type_id", "itinerary_id", "bus_route_id",
        "dep_time", "dep_time_sec", "dep_stop_id", "dep_stop_name", "arr_time",
        "arr_time_sec", "arr_stop_id", "arr_stop_name", "trip_type", "duration",
        "distance", "height", "num_bus_stops", "bool_service",
    }
    if not trips or not trip_fields.issubset(trips[0]):
        raise ValueError("Unexpected trip_set.xlsx schema")
    stop_ids = {int(s["id"]) for s in stops}
    itinerary_ids = {int(i["id"]) for i in itineraries}
    route_ids = {int(r["bus_route_id"]) for r in routes}
    ids = [int(t["trip_id"]) for t in trips]
    if len(set(ids)) != len(ids):
        raise ValueError("Duplicate trip IDs in fixed service set")
    source_trips = []
    for excel_row, row in enumerate(trips, start=2):
        dep = parse_day_time(row["dep_time"])
        arr = parse_day_time(row["arr_time"])
        duration = parse_duration(row["duration"])
        if dep != int(row["dep_time_sec"]) or arr != int(row["arr_time_sec"]):
            raise ValueError(f"Source string/seconds mismatch at trip_set physical row {excel_row}")
        if arr - dep != duration:
            raise ValueError(f"Source duration does not reconcile at trip_set physical row {excel_row}")
        if dep % 30 or arr % 30 or duration % 30:
            raise ValueError("Source trip time cannot be represented at exact half-minute resolution")
        if int(row["trip_type"]) != 1 or row["bool_service"] is not True:
            raise ValueError(f"{config['name']} raw timetable contains an unexpected nonservice row")
        if int(row["itinerary_id"]) not in itinerary_ids or int(row["bus_route_id"]) not in route_ids:
            raise ValueError("Unresolved source itinerary or route reference")
        if int(row["dep_stop_id"]) not in stop_ids or int(row["arr_stop_id"]) not in stop_ids:
            raise ValueError("Unresolved source trip endpoint")
        source_row = {k: row[k] for k in sorted(trip_fields)}
        source_row.update({
            "source_excel_sheet": "Sheet1",
            "source_excel_physical_row": excel_row,
            "source_time_unit": "seconds",
            "source_distance_unit": "km",
            "source_height_difference_unit": "m",
            "source_energy_kwh": None,
            "source_energy_status": "not supplied in raw service timetable",
            "_dep_time_sec_exact": dep,
            "_arr_time_sec_exact": arr,
            "_duration_sec_exact": duration,
            "_native_id": f"{config['trip_prefix']}{int(row['trip_id']):04d}",
        })
        source_trips.append(source_row)
    if len({int(x["trip_id"]) for x in source_trips}) != config["service_count"]:
        raise ValueError(f"Service set did not retain all {config['service_count']} unique source trips")

    mat_data = loadmat(io.BytesIO(zf.read(prefix + "deadhead_trip_matrix.mat")),
                       squeeze_me=True, struct_as_record=False)
    public_fields = [v for k, v in mat_data.items() if not k.startswith("__")]
    if len(public_fields) != 1 or not hasattr(public_fields[0], "_fieldnames"):
        raise ValueError("Unexpected deadhead MATLAB data structure")
    matrix = public_fields[0]
    matrix_stop_ids = [int(v) for v in np.asarray(matrix.bus_stop_id).reshape(-1)]
    distances_m = np.asarray(matrix.distances, dtype=float)
    times_s = np.asarray(matrix.times, dtype=float)
    heights_m = np.asarray(matrix.height_difference, dtype=float)
    n = len(matrix_stop_ids)
    if n != len(stops) or any(a.shape != (n, n) for a in (distances_m, times_s, heights_m)):
        raise ValueError("Deadhead matrix is not a complete stop-aligned square matrix")
    if not all(np.isfinite(a).all() for a in (distances_m, times_s, heights_m)):
        raise ValueError("Missing/unknown deadhead edge; adapter will not impute")
    if (distances_m < 0).any() or (times_s < 0).any():
        raise ValueError("Negative source deadhead distance or time")
    if any(int(x) != x for x in times_s.reshape(-1)):
        raise ValueError("Source deadhead time is not an integral second")
    if any(int(x) % 30 for x in times_s.reshape(-1)):
        raise ValueError("Source deadhead time is outside the exact half-minute model grid")
    if set(matrix_stop_ids) != stop_ids:
        raise ValueError("Deadhead matrix stop IDs differ from source stop workbook")
    matrix_index = {stop_id: i for i, stop_id in enumerate(matrix_stop_ids)}
    dhd_arcs = []
    for oi, origin in enumerate(matrix_stop_ids):
        for di, destination in enumerate(matrix_stop_ids):
            dhd_arcs.append({
                "origin_stop_id": origin,
                "destination_stop_id": destination,
                "distance_m_source": float(distances_m[oi, di]),
                "time_sec_source": int(times_s[oi, di]),
                "height_difference_m_source": float(heights_m[oi, di]),
                "source_energy_kwh": None,
                "source_energy_status": "not supplied; modeled separately from Table 2 assumptions",
            })
    depot_ids = sorted(int(s["id"]) for s in stops if s["b_depot"] in (True, 1))
    if tuple(depot_ids) != config["depot_ids"]:
        raise ValueError(f"Expected flagged depots {config['depot_ids']}; found {depot_ids}")
    counts_by_route: dict[int, int] = {}
    for t in source_trips:
        rid = int(t["bus_route_id"])
        counts_by_route[rid] = counts_by_route.get(rid, 0) + 1
    declared_by_route = {int(r["bus_route_id"]): int(r["num_trips"]) for r in routes}
    if counts_by_route != declared_by_route:
        raise ValueError("Trip route totals do not match source route metadata")

    return {
        "operator_directory": config["directory"],
        "operator_name": config["name"],
        "service_trips": source_trips,
        "stops": [
            {k: s[k] for k in sorted(s)} | {
                "source_excel_sheet": "Sheet1",
                "source_excel_physical_row": i + 2,
                "coordinates_unit": "latitude/longitude WGS84 degrees",
                "height_unit": "m",
                "place_id_native": f"P{int(s['id'])}",
            }
            for i, s in enumerate(stops)
        ],
        "routes": [
            {k: r[k] for k in sorted(r)} | {
                "source_excel_sheet": "Sheet1",
                "source_excel_physical_row": i + 2,
                "service_mileage_unit": "km",
            }
            for i, r in enumerate(routes)
        ],
        "itineraries": [
            {k: it[k] for k in sorted(it)} | {
                "source_excel_sheet": "Sheet1",
                "source_excel_physical_row": i + 2,
                "duration_unit": "min",
                "distance_unit": "km",
                "height_difference_unit": "m",
            }
            for i, it in enumerate(itineraries)
        ],
        "deadhead_stop_order_source": matrix_stop_ids,
        "deadhead_arcs_directed_row_major": dhd_arcs,
        "source_day_type_ids": sorted({int(t["day_type_id"]) for t in source_trips}),
        "calendar_date_present": False,
        "source_member_manifest": member_provenance(zf, archive_path, operator),
        "matrix_index_by_stop_id": matrix_index,
    }


def _place(stop_id: int) -> str:
    return f"P{int(stop_id)}"


def _arc(source: dict, i: int, j: int) -> dict:
    n = len(source["deadhead_stop_order_source"])
    return source["deadhead_arcs_directed_row_major"][i * n + j]


def _arc_between(source: dict, origin: int, destination: int) -> dict:
    ix = source["matrix_index_by_stop_id"]
    return _arc(source, ix[str(origin)] if str(origin) in ix else ix[origin],
                ix[str(destination)] if str(destination) in ix else ix[destination])


def _modeled_dhd_energy(distance_m: float, time_sec: int) -> dict:
    km = distance_m / 1000.0
    traction = EB3["deadhead_traction_kwh_per_km"] * km
    auxiliary = EB3["auxiliary_kwh_per_hour"] * time_sec / 3600.0
    return {"traction_kwh": traction, "auxiliary_kwh": auxiliary,
            "total_kwh": traction + auxiliary,
            "formula": "0.96 kWh/km * source DHD distance + 16 kWh/h * source DHD travel time"}


def _modeled_service_energy(source_trip: dict) -> dict:
    km = float(source_trip["distance"])
    duration_h = source_trip["_duration_sec_exact"] / 3600.0
    traction = EB3["service_traction_kwh_per_km"] * km
    auxiliary = EB3["auxiliary_kwh_per_hour"] * duration_h
    return {"traction_kwh": traction, "auxiliary_kwh": auxiliary,
            "total_kwh": traction + auxiliary,
            "formula": "1.20 kWh/km * source service distance + 16 kWh/h * source service duration"}


def build_depot_variant(source: dict, depot_id: int,
                        operator: str = "hildenbrand") -> tuple[object, dict]:
    """Build one explicit single-depot NativeCase plus movement provenance."""
    config = OPERATOR_CONFIGS[operator]
    if depot_id not in {int(s["id"]) for s in source["stops"] if s["b_depot"] in (True, 1)}:
        raise ValueError(f"Depot {depot_id} is not source-flagged")
    from egglab import native_recharge as nr

    source_trips = sorted(source["service_trips"],
                          key=lambda t: (t["_dep_time_sec_exact"], t["_arr_time_sec_exact"], t["_native_id"]))
    source_depot_ids = {int(s["id"]) for s in source["stops"] if s["b_depot"] in (True, 1)}
    service_native = []
    trip_lookup = {}
    energy_by_trip = {}
    for t in source_trips:
        start_s = int(t["_dep_time_sec_exact"])
        end_s = int(t["_arr_time_sec_exact"])
        e = _modeled_service_energy(t)
        if e["total_kwh"] > EB3["usable_battery_kwh"]:
            raise ValueError(f"Single service {t['_native_id']} exceeds usable battery")
        native = nr.Trip(
            id=t["_native_id"],
            start_min=exact_native_minutes(start_s),
            end_min=exact_native_minutes(end_s),
            start_place=_place(int(t["dep_stop_id"])),
            end_place=_place(int(t["arr_stop_id"])),
            energy_kwh=e["total_kwh"],
        )
        service_native.append(native)
        trip_lookup[native.id] = t
        energy_by_trip[native.id] = {
            "source_energy_kwh": None,
            "modeled_energy_kwh": e["total_kwh"],
            "energy_components": e,
            "source_member": config["member_prefix"] + "trip_set.xlsx",
            "source_excel_physical_row": t["source_excel_physical_row"],
        }

    stop_order = source["deadhead_stop_order_source"]
    ix = {int(v): i for i, v in enumerate(stop_order)}
    moves = []
    provenance = []
    pair_counts = {"direct": 0, "depot": 0, "pullout": 0, "pullin": 0}

    def add_leg(origin_id: int, destination_id: int, depart_s: int, arrive_s: int,
                role: str, source_arc: dict):
        if arrive_s < depart_s:
            raise ValueError("Negative-duration movement leg")
        if arrive_s - depart_s != source_arc["time_sec_source"] and role != "off_depot_stationary_wait":
            raise ValueError("Movement leg duration does not equal the exact source DHD time")
        energy = ({"traction_kwh": 0.0,
                   "auxiliary_kwh": EB3["auxiliary_kwh_per_hour"] * (arrive_s-depart_s) / 3600,
                   "total_kwh": EB3["auxiliary_kwh_per_hour"] * (arrive_s-depart_s) / 3600,
                   "formula": "16 kWh/h * explicit off-depot stationary wait; zero traction"}
                  if role == "off_depot_stationary_wait"
                  else _modeled_dhd_energy(source_arc["distance_m_source"], source_arc["time_sec_source"]))
        leg = nr.Leg(_place(origin_id), _place(destination_id),
                     exact_native_minutes(depart_s), exact_native_minutes(arrive_s), energy["total_kwh"])
        return leg, {
            "origin_stop_id": origin_id,
            "destination_stop_id": destination_id,
            "depart_sec": depart_s,
            "arrive_sec": arrive_s,
            "duration_sec": arrive_s - depart_s,
            "source_distance_m": source_arc["distance_m_source"] if role != "off_depot_stationary_wait" else 0.0,
            "source_travel_time_sec": source_arc["time_sec_source"] if role != "off_depot_stationary_wait" else 0,
            "source_height_difference_m": source_arc["height_difference_m_source"] if role != "off_depot_stationary_wait" else 0.0,
            "source_energy_kwh": None,
            "modeled_energy_kwh": energy["total_kwh"],
            "modeled_energy_components": energy,
            "leg_role": role,
            "source_member": (config["member_prefix"] + "deadhead_trip_matrix.mat"
                              if role != "off_depot_stationary_wait" else None),
            "energy_status": "modeled from declared EB-3 assumptions; no source leg energy supplied",
        }

    # Pull-out: leave the selected depot as late as possible while arriving
    # exactly at the fixed service start. There is no decision variable for
    # departure time in this enumerated mode.
    for t in source_trips:
        tid = t["_native_id"]
        dest = int(t["dep_stop_id"])
        a = _arc_between(source, depot_id, dest)
        travel_s = a["time_sec_source"]
        depart_s = int(t["_dep_time_sec_exact"]) - travel_s
        arrive_s = int(t["_dep_time_sec_exact"])
        if HORIZON_START_SEC <= depart_s <= arrive_s <= HORIZON_END_SEC:
            leg, leg_meta = add_leg(depot_id, dest, depart_s, arrive_s, "deadhead_pullout", a)
            mid = f"pullout_D{depot_id}_{tid}"
            moves.append(nr.Movement(mid, "pullout", None, tid, (leg,), None))
            provenance.append({"movement_id": mid, "kind": "pullout", "before": None, "after": tid,
                               "departure_policy": "latest feasible: service_start minus directed source DHD time",
                               "legs": [leg_meta]})
            pair_counts["pullout"] += 1

    # Inter-service connection modes use each ordered service pair and retain
    # direction-specific source travel. Direct moves begin immediately after
    # service; they end with an explicit off-depot stationary wait (if any).
    for before in source_trips:
        bid = before["_native_id"]
        from_stop = int(before["arr_stop_id"])
        service_end = int(before["_arr_time_sec_exact"])
        for after in source_trips:
            aid = after["_native_id"]
            if bid == aid:
                continue
            to_stop = int(after["dep_stop_id"])
            service_start = int(after["_dep_time_sec_exact"])
            # Direct mode: immediate DHD travel followed by same-place wait.
            a = _arc_between(source, from_stop, to_stop)
            arrival_s = service_end + a["time_sec_source"]
            if arrival_s <= service_start and arrival_s <= HORIZON_END_SEC:
                legs, leg_meta = [], []
                travel_leg, travel_meta = add_leg(from_stop, to_stop, service_end, arrival_s,
                                                  "deadhead_direct", a)
                legs.append(travel_leg); leg_meta.append(travel_meta)
                if arrival_s < service_start:
                    if to_stop in source_depot_ids:
                        raise ValueError(
                            "Direct-mode stationary wait is at a source-flagged depot; "
                            "the selected policy sets depot auxiliary off and requires explicit handling"
                        )
                    wait_arc = {"distance_m_source": 0.0, "time_sec_source": 0,
                                "height_difference_m_source": 0.0}
                    wait_leg, wait_meta = add_leg(to_stop, to_stop, arrival_s, service_start,
                                                  "off_depot_stationary_wait", wait_arc)
                    legs.append(wait_leg); leg_meta.append(wait_meta)
                mid = f"direct_{bid}_{aid}"
                moves.append(nr.Movement(mid, "direct", bid, aid, tuple(legs), None))
                provenance.append({"movement_id": mid, "kind": "direct", "before": bid, "after": aid,
                                   "departure_policy": "immediate after prior service; any residual time is fixed destination wait",
                                   "legs": leg_meta})
                pair_counts["direct"] += 1

            # Depot detour: inbound immediately, outbound latest feasible to
            # arrive exactly at the next mandatory service. This maximizes the
            # explicitly modeled depot dwell/recharge window; no depot idle
            # auxiliary load is assumed. It is one fixed mode per pair.
            in_arc = _arc_between(source, from_stop, depot_id)
            out_arc = _arc_between(source, depot_id, to_stop)
            in_arrive_s = service_end + in_arc["time_sec_source"]
            out_depart_s = service_start - out_arc["time_sec_source"]
            if (in_arrive_s <= out_depart_s and out_depart_s <= service_start
                    and HORIZON_START_SEC <= service_end <= in_arrive_s <= HORIZON_END_SEC):
                in_leg, in_meta = add_leg(from_stop, depot_id, service_end, in_arrive_s,
                                          "deadhead_depot_inbound", in_arc)
                out_leg, out_meta = add_leg(depot_id, to_stop, out_depart_s, service_start,
                                            "deadhead_depot_outbound", out_arc)
                mid = f"depot_D{depot_id}_{bid}_{aid}"
                moves.append(nr.Movement(mid, "depot", bid, aid, (in_leg, out_leg), 1))
                provenance.append({"movement_id": mid, "kind": "depot", "before": bid, "after": aid,
                                   "depot_id": depot_id,
                                   "depot_split_after_leg": 1,
                                   "departure_policy": "inbound immediately after service; outbound latest feasible to arrive at next service start",
                                   "depot_dwell_start_sec": in_arrive_s,
                                   "depot_dwell_end_sec": out_depart_s,
                                   "depot_auxiliary_load_assumption": "off while parked at depot",
                                   "legs": [in_meta, out_meta]})
                pair_counts["depot"] += 1

    # Pull-in: leave immediately after the final fixed service event; the
    # return trip and terminal full-charge window are then explicit.
    for t in source_trips:
        tid = t["_native_id"]
        origin = int(t["arr_stop_id"])
        a = _arc_between(source, origin, depot_id)
        depart_s = int(t["_arr_time_sec_exact"])
        arrive_s = depart_s + a["time_sec_source"]
        if HORIZON_START_SEC <= depart_s <= arrive_s <= HORIZON_END_SEC:
            leg, leg_meta = add_leg(origin, depot_id, depart_s, arrive_s, "deadhead_pullin", a)
            mid = f"pullin_{tid}_D{depot_id}"
            moves.append(nr.Movement(mid, "pullin", tid, None, (leg,), None))
            provenance.append({"movement_id": mid, "kind": "pullin", "before": tid, "after": None,
                               "departure_policy": "immediate after service; arrival defines terminal depot availability",
                               "terminal_charge_window_sec": [arrive_s, HORIZON_END_SEC],
                               "legs": [leg_meta]})
            pair_counts["pullin"] += 1

    # Complete the direct time-of-day market at exact hourly minute boundaries.
    market_edges = tuple(range(0, 30 * 60 + 1, 60))
    case_name = f"{config['case_prefix']}_{depot_id}_eb3"
    case = nr.NativeCase(
        name=case_name,
        trips=tuple(service_native),
        movements=tuple(moves),
        resources=(nr.Resource(0, 1800, 360.0, 360.0, 1),),
        market_edges_min=market_edges,
        depot=_place(depot_id),
        max_vehicles=config["max_vehicles"],
        battery_kwh=400.0,
        reserve_kwh=0.0,
        terminal_open_min=0,
        recharge_deadline_min=1800,
        vehicle_cost=VEHICLE_COST_SYNTHETIC,
        deadhead_cost_per_min=DEADHEAD_COST_PER_MIN_SYNTHETIC,
        efficiency=1.0,
        graph_scope="declared-movement-modes-only",
    )
    if pair_counts["pullout"] != config["service_count"] or pair_counts["pullin"] != config["service_count"]:
        raise ValueError("Every mandatory trip must have a source-resolved pullout and pullin in this depot variant")
    model_trips = []
    for nt in case.trips:
        t = trip_lookup[nt.id]
        model_trips.append({
            "native": asdict(nt),
            "source_trip_id": int(t["trip_id"]),
            "source_excel_physical_row": t["source_excel_physical_row"],
            "source_times_sec": {"start": t["_dep_time_sec_exact"], "end": t["_arr_time_sec_exact"]},
            "source_distance_km": float(t["distance"]),
            "source_energy_kwh": None,
            "model_energy": energy_by_trip[nt.id],
        })
    return case, {
        "case_name": case.name,
        "case_identity": case.identity(),
        "variant_policy": ("single source-flagged depot only; not the source two-depot fleet model"
                           if operator == "hildenbrand" else
                           "single source-flagged depot EGG scenario; not a publisher fleet result"),
        "selected_depot_id": depot_id,
        "selected_depot_native_place": _place(depot_id),
        "service_count": len(case.trips),
        "movement_mode_counts": pair_counts,
        "direct_off_depot_wait_mode_count": sum(
            1 for row in provenance if row["kind"] == "direct"
            for leg in row["legs"] if leg["leg_role"] == "off_depot_stationary_wait"
        ),
        "direct_wait_at_flagged_depot_count": sum(
            1 for row in provenance if row["kind"] == "direct"
            for leg in row["legs"] if leg["leg_role"] == "off_depot_stationary_wait"
            and leg["origin_stop_id"] in source_depot_ids
        ),
        "movement_modes": [asdict(m) for m in case.movements],
        "movement_provenance": provenance,
        "trips": model_trips,
        "resource_policy": asdict(case.resources[0]),
        "market_edges_min": list(case.market_edges_min),
        "vehicle_cap": config["max_vehicles"],
        "initial_inventory_policy": "every used bus starts at full usable EB-3 inventory (400 kWh)",
        "terminal_policy": "every used bus returns to its selected depot and must restore full usable inventory by 30:00 (1800 min); terminal opening is 00:00, so all post-pull-in time is available",
        "availability_horizon_sec": [HORIZON_START_SEC, HORIZON_END_SEC],
        "time_representation": "source data remains integer seconds; every native minute value is an exact integer or half-minute conversion; no rounding",
        "energy_model": {
            "paper_concept_id": EB3["paper_vehicle_id"],
            "paper_concept_name": EB3["paper_vehicle_name"],
            "paper_source_table": "Sistig et al. (2025), Table 2; DOI 10.1038/s44333-025-00030-y",
            "source_energy_for_services_or_deadheads": None,
            "service_traction_kwh_per_km": EB3["service_traction_kwh_per_km"],
            "deadhead_traction_kwh_per_km": EB3["deadhead_traction_kwh_per_km"],
            "auxiliary_kwh_per_hour": EB3["auxiliary_kwh_per_hour"],
            "service_energy_formula": "1.20 * source distance km + 16 * source service duration hours",
            "movement_energy_formula": "0.96 * source deadhead distance km + 16 * source deadhead travel hours",
            "off_depot_stationary_wait_formula": "16 * explicit stationary off-depot wait hours; added as same-place direct-mode leg",
            "depot_parked_auxiliary": "assumed off per research lead's selected EGG protocol",
            "energy_values_are_modelled_not_observed": True,
        },
        "charging_model": {
            "usable_battery_kwh": 400.0,
            "installed_battery_kwh_reference_only": 600.0,
            "reserve_kwh": 0.0,
            "efficiency": 1.0,
            "per_bus_kw": 360.0,
            "shared_grid_kw": 360.0,
            "connectors": 1,
            "relationship_to_publisher": ("EGG idealization, not a reproduction of publisher's assumed one charger with four outputs"
                                          if operator == "hildenbrand" else
                                          "EGG idealization; charger configuration is not observed in Eberbach source inputs"),
        },
        "cost_policy": {
            "vehicle_cost": VEHICLE_COST_SYNTHETIC,
            "vehicle_cost_unit": "f100 synthetic currency per used bus",
            "deadhead_cost_per_min": DEADHEAD_COST_PER_MIN_SYNTHETIC,
            "not_source_economic_data": True,
        },
        "publisher_schedule_or_charging_output_used": False,
        "scientific_campaign_frozen": False,
    }


def _construct_and_replay_one_trip_witness(case, compiled: dict) -> dict:
    """Construct a transparent one-service-per-bus witness; never optimize."""
    from egglab import native_recharge as nr

    pullouts = {m.after: m for m in case.movements if m.kind == "pullout"}
    pullins = {m.before: m for m in case.movements if m.kind == "pullin"}
    if set(pullouts) != {t.id for t in case.trips} or set(pullins) != {t.id for t in case.trips}:
        raise ValueError("Cannot construct one-trip-per-bus witness: incomplete pullout/pullin coverage")

    paths = []
    for trip in case.trips:
        first, last = pullouts[trip.id], pullins[trip.id]
        path_energy = trip.energy_kwh + sum(x.energy_kwh for x in first.legs + last.legs)
        if path_energy > case.battery_kwh - case.reserve_kwh + nr.ENERGY_TOL:
            raise ValueError(f"One-trip bus {trip.id} exceeds usable inventory under fixed movement modes")
        release_min = last.legs[-1].arrive_min
        paths.append({"trip": trip, "pullout": first, "pullin": last,
                      "energy_kwh": path_energy, "release_min": release_min})

    # All terminal jobs share one deadline, so release-ordered contiguous
    # sessions are a deterministic single-connector witness if their finish
    # times fit. This is a constructive replay certificate, never an optimum.
    paths.sort(key=lambda x: (x["release_min"], x["trip"].id))
    cursor = float(case.terminal_open_min)
    vehicles, charges = [], []
    for vehicle_id, path in enumerate(paths):
        start = max(float(path["release_min"]), cursor)
        duration_min = path["energy_kwh"] * 60.0 / case.resources[0].per_bus_kw
        end = start + duration_min
        if end > case.recharge_deadline_min + nr.TIME_TOL_MIN:
            raise ValueError("One-trip-per-bus single-connector witness misses full-return deadline")
        vehicles.append({"vehicle": vehicle_id, "trips": [path["trip"].id],
                         "movements": [path["pullout"].id, path["pullin"].id]})
        charges.append({"vehicle": vehicle_id, "movement": path["pullin"].id,
                        "connector": 0, "start_min": start, "end_min": end,
                        "grid_kwh": path["energy_kwh"]})
        cursor = end

    load = [0.0] * (len(case.market_edges_min) - 1)
    for charge in charges:
        duration = charge["end_min"] - charge["start_min"]
        for period, (lo, hi) in enumerate(zip(case.market_edges_min, case.market_edges_min[1:])):
            overlap = max(0.0, min(charge["end_min"], hi) - max(charge["start_min"], lo))
            if overlap:
                load[period] += charge["grid_kwh"] * overlap / duration
    plan = {"schema": nr.SCHEMA, "case_identity": case.identity(),
            "vehicles": vehicles, "charges": charges, "load": load,
            "ops_cost": len(vehicles) * case.vehicle_cost}
    replay = nr.replay_native(case, plan)

    service_energy = math.fsum(t.energy_kwh for t in case.trips)
    movement_energy = math.fsum(
        leg.energy_kwh for path in paths
        for movement in (path["pullout"], path["pullin"]) for leg in movement.legs
    )
    eligible_mode_interval_pairs = sum(
        len(interval["visits"]) for interval in compiled["intervals"] if interval["rate_kw"] > 0
    )
    return {
        "classification": "constructed one-trip-per-bus feasible witness; not optimized",
        "replay_ok": replay["replay_ok"],
        "used_vehicle_count": len(vehicles),
        "service_energy_total_kwh": service_energy,
        "one_trip_per_bus_pullout_pullin_energy_total_kwh": movement_energy,
        "one_trip_per_bus_total_battery_energy_kwh": service_energy + movement_energy,
        "maximum_one_trip_bus_energy_kwh": max(p["energy_kwh"] for p in paths),
        "latest_pullin_arrival_min": max(p["release_min"] for p in paths),
        "last_terminal_charge_end_min": max(c["end_min"] for c in charges),
        "total_terminal_charge_kwh": replay["grid_kwh"],
        "maximum_shared_grid_kw": replay["max_grid_kw"],
        "compiled_resource_interval_count": len(compiled["intervals"]),
        "eligible_mode_interval_pairs_per_vehicle": eligible_mode_interval_pairs,
        "candidate_charge_variables_at_max_vehicles": eligible_mode_interval_pairs * case.max_vehicles,
        "solver_invoked": False,
    }


def estimate_compact_model_dimensions(case, compiled: dict) -> dict:
    """Count path-flow variables/rows from compiled intervals; allocate no MIP."""
    n = len(case.trips)
    m = len(case.movements)
    intervals = compiled["intervals"]
    charge = sum(len(x["visits"]) for x in intervals if x["rate_kw"] > 0)
    depot = sum(x.kind == "depot" for x in case.movements)
    pullin = sum(x.kind == "pullin" for x in case.movements)
    periods = len(case.market_edges_min) - 1
    variables = m + 2*n + charge + periods
    rows = (2 + 3*n + charge + len(intervals) + 2*m + 2*depot
            + pullin + 2 + periods)
    return {
        "formulation": "egg-native-pathflow-v3-energy-band-orphan-projection",
        "basis": "combinatorial count of native_pathflow.build_feasible_model; no model allocated",
        "movement_binary_variables": m,
        "service_soc_continuous_variables": 2*n,
        "charge_interval_continuous_variables": charge,
        "market_load_continuous_variables": periods,
        "total_variables_before_objective": variables,
        "estimated_constraints_before_objective": rows,
        "compiled_resource_intervals": len(intervals),
        "ordered_service_pairs": n*(n-1),
        "solver_invoked": False,
        "model_allocated": False,
    }


def json_ready_source(source: dict) -> dict:
    """Remove only private in-script indexes and parsing scratch fields."""
    value = json.loads(json.dumps(source, allow_nan=False))
    value.pop("matrix_index_by_stop_id", None)
    for row in value["service_trips"]:
        for key in ("_dep_time_sec_exact", "_arr_time_sec_exact", "_duration_sec_exact", "_native_id"):
            row.pop(key, None)
    return value


def native_case_from_payload(variant: dict):
    """Rebuild a NativeCase from the shipped JSON without reading the source ZIP."""
    from egglab import native_recharge as nr

    trips = tuple(nr.Trip(**t["native"]) for t in variant["trips"])
    movements = []
    for m in variant["movement_modes"]:
        legs = tuple(nr.Leg(**leg) for leg in m["legs"])
        movements.append(nr.Movement(
            id=m["id"], kind=m["kind"], before=m["before"], after=m["after"],
            legs=legs, depot_split=m["depot_split"],
        ))
    charging = variant["charging_model"]
    cost = variant["cost_policy"]
    return nr.NativeCase(
        name=variant["case_name"], trips=trips, movements=tuple(movements),
        resources=(nr.Resource(**variant["resource_policy"]),),
        market_edges_min=tuple(variant["market_edges_min"]), depot=variant["selected_depot_native_place"],
        max_vehicles=int(variant["vehicle_cap"]), battery_kwh=charging["usable_battery_kwh"],
        reserve_kwh=charging["reserve_kwh"], terminal_open_min=0,
        recharge_deadline_min=1800, vehicle_cost=cost["vehicle_cost"],
        deadhead_cost_per_min=cost["deadhead_cost_per_min"], efficiency=charging["efficiency"],
        graph_scope="declared-movement-modes-only",
    )


def build_document(archive_path: Path, validate_native: bool = False,
                   operator: str = "hildenbrand") -> dict:
    config = OPERATOR_CONFIGS[operator]
    if archive_path.stat().st_size != ZIP_BYTES:
        raise ValueError(f"Unexpected archive size: {archive_path.stat().st_size}")
    if sha256_file(archive_path) != ZIP_SHA256 or md5_file(archive_path) != ZIP_MD5:
        raise ValueError("Archive bytes do not match the DOI-pinned publisher checksums")
    with zipfile.ZipFile(archive_path) as zf:
        source = extract_source(zf, archive_path, operator)
    depot_ids = sorted(int(s["id"]) for s in source["stops"] if s["b_depot"] in (True, 1))
    if operator == "eberbach" and len(source["service_trips"]) * (len(source["service_trips"])-1) > 12_000:
        raise ValueError("Source-only size screen exceeds the bounded Eberbach intake")
    variants = []
    for depot_id in depot_ids:
        case, variant = build_depot_variant(source, depot_id, operator)
        if validate_native:
            from egglab import native_recharge as nr
            nr.validate_case(case)
            compiled = nr.compile_case(case)  # schema compilation only; no solver
            variant["native_schema_validation"] = "passed"
            variant["compiled_resource_interval_count"] = len(compiled["intervals"])
            if operator == "eberbach":
                variant["compact_preflight_dimensions"] = estimate_compact_model_dimensions(case, compiled)
                try:
                    variant["preflight_witness_and_dimensions"] = _construct_and_replay_one_trip_witness(case, compiled)
                except ValueError as exc:
                    variant["preflight_witness_and_dimensions"] = {
                        "classification": "conservative one-trip-per-bus construction failed; not a fleet infeasibility proof",
                        "reason": str(exc), "solver_invoked": False,
                    }
            else:
                variant["preflight_witness_and_dimensions"] = _construct_and_replay_one_trip_witness(case, compiled)
        else:
            variant["native_schema_validation"] = "not-run (use --validate-native after exact half-minute support is present)"
        variants.append(variant)
    source_safe = json_ready_source(source)
    latest_service = max(source["service_trips"], key=lambda t: t["_arr_time_sec_exact"])
    return {
        "classification": "CC BY 4.0 derived public microcase; contains public timetable/movement rows and explicitly modeled energy; do not include source archive bytes",
        "dataset": {
            "title": "Dataset for Evaluating Costs and Operations of Public Bus Fleet Electrification",
            "doi": DOI,
            "figshare_url": FIGSHARE_URL,
            "license": "CC BY 4.0",
            "license_url": LICENSE_URL,
            "dataset_citation": "Sistig, Hubert Maximilian; Sinhuber, Philipp; Rogge, Matthias; Sauer, D.U. (2025). Dataset for Evaluating Costs and Operations of Public Bus Fleet Electrification. figshare. Dataset. https://doi.org/10.6084/m9.figshare.26088190.v1",
            "paper_citation": "Sistig, H. M., Sinhuber, P., Rogge, M., & Sauer, D. U. (2025). Evaluating costs and operations of public bus fleet electrification. npj Sustainable Mobility and Transport 2, 15. https://doi.org/10.1038/s44333-025-00030-y",
            "attribution": "Contains adapted data from Sistig et al. (2025), Dataset for Evaluating Costs and Operations of Public Bus Fleet Electrification, CC BY 4.0. Source values remain labeled; all native modes and EB-3 energy fields are a new EGG transformation. See this file and the protocol for changes and assumptions.",
        },
        "source_archive_provenance": source["source_member_manifest"],
        "source_field_semantics": {
            "trip_set.dep_time/arr_time": "dd:HH:MM:SS, converted losslessly to integer seconds since selected-day boundary",
            "trip_set.dep_time_sec/arr_time_sec": "source seconds; checked against timestamp text",
            "trip_set.duration": "HH:MM:SS; source duration checked against arrival minus departure",
            "trip_set.distance": "km",
            "trip_set.height": "m height difference",
            "deadhead_trip_matrix.distances": "m; directional matrix values retained exactly",
            "deadhead_trip_matrix.times": "s; directional matrix values retained exactly",
            "deadhead_trip_matrix.height_difference": "m",
            "source_service_or_deadhead_energy": "not provided",
        },
        "source_data": source_safe,
        "readiness_summary": {
            "mandatory_service_count": len(source["service_trips"]),
            "last_service_arrival": {
                "trip_id": int(latest_service["trip_id"]),
                "timestamp_sec_from_selected_day_boundary": latest_service["_arr_time_sec_exact"],
                "timestamp_min": exact_native_minutes(latest_service["_arr_time_sec_exact"]),
            },
            "deadline_sec_from_selected_day_boundary": HORIZON_END_SEC,
            "deadline_min": exact_native_minutes(HORIZON_END_SEC),
            "units_note": "Source event values remain integer seconds; native timestamps are exact integer/half-minute values.",
            "preflight_only": True,
            "solver_invoked": False,
        },
        "modeling_assumptions": {
            "vehicle_concept": EB3,
            "single_depot_variants": depot_ids,
            "time_horizon": {"start_sec": HORIZON_START_SEC, "deadline_sec": HORIZON_END_SEC,
                             "deadline_label": "30:00 (30 hours after selected-day boundary)"},
            "charging": {"usable_battery_kwh": 400.0, "reserve_kwh": 0.0, "efficiency": 1.0,
                         "per_bus_kw": 360.0, "shared_grid_kw": 360.0, "connectors": 1,
                         "policy": "EGG idealization; one connector, constant grid-side 360 kW"},
            "costs": {"vehicle_cost": 100.0, "unit": "f100 synthetic currency per used bus",
                      "deadhead_cost_per_min": 0.0, "source_economic_data": False},
            "terminal_open_sec": 0,
            "max_used_vehicles": config["max_vehicles"],
            "departure_policies": {
                "pullout": "latest feasible: leave selected depot at service start minus directed DHD travel time",
                "direct": "leave immediately after prior service; add explicit off-depot stationary wait at next service origin",
                "depot_detour": "inbound immediately after service, outbound latest feasible to arrive at next service start; fixes the depot charging window",
                "pullin": "leave immediately after final service; recharge at selected depot through the 30:00 deadline",
            },
            "depot_idle_auxiliary": "off while parked at depot",
            "direct_wait_policy": "auxiliary applies only to off-depot wait; adapter fails closed if a direct wait occurs at a source-flagged depot",
            "source_energy_observed": False,
            "publisher_derived_schedule_or_charging_used": False,
            "scientific_campaign_frozen": False,
        },
        "native_cases": variants,
    }


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--archive", type=Path, default=DEFAULT_ARCHIVE)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--operator", choices=tuple(OPERATOR_CONFIGS), default="hildenbrand")
    parser.add_argument("--validate-native", action="store_true",
                        help="run pure NativeCase validation/compile (requires exact half-minute time support; never solves)")
    args = parser.parse_args(argv)
    if args.output is None:
        args.output = DEFAULT_OUTPUT if args.operator == "hildenbrand" else EBERBACH_OUTPUT
    document = build_document(args.archive, validate_native=args.validate_native,
                              operator=args.operator)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(document, indent=2, ensure_ascii=False, allow_nan=False) + "\n", encoding="utf-8")
    print(json.dumps({
        "output": str(args.output),
        "archive_sha256": document["source_archive_provenance"]["archive_sha256"],
        "service_count": len(document["source_data"]["service_trips"]),
        "directed_matrix_arc_count": len(document["source_data"]["deadhead_arcs_directed_row_major"]),
        "variants": [{"case_name": v["case_name"], "depot": v["selected_depot_id"],
                      "mode_counts": v["movement_mode_counts"],
                      "validation": v["native_schema_validation"]} for v in document["native_cases"]],
        "solver_invoked": False,
    }, indent=2))


if __name__ == "__main__":
    main()
